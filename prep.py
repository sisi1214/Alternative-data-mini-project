"""Prep: C&SD jewellery monthly -> tidy CSV -> fiscal-quarter yoy; Luk Fook SSSG tidy; gold quarterly; merged panel.
Fiscal year = April-March. Quarter label NQyyyy = Nth quarter of the FY STARTING in April of yyyy (your sheet's convention):
1Q = Apr-Jun, 2Q = Jul-Sep, 3Q = Oct-Dec, 4Q = Jan-Mar(yyyy+1)."""
import openpyxl, pandas as pd, numpy as np, re, os
OUT = '/mnt/user-data/outputs'; os.makedirs(OUT, exist_ok=True)
CSD = '/mnt/user-data/uploads/Jewellery__watches_and_clocks__and_valuable_gifts_2022-2026.xlsx'
LF = '/mnt/user-data/uploads/Luk_Fook_analysis.xlsx'

# ---------- 1. C&SD monthly tidy ----------
ws = openpyxl.load_workbook(CSD, data_only=True).active
yr, cur = {}, None
for c in range(2, ws.max_column + 1):
    v = ws.cell(1, c).value
    cur = v if v is not None else cur
    yr[c] = cur
rows = []
for c0 in range(2, ws.max_column + 1, 3):
    m = ws.cell(2, c0).value
    rows.append(dict(year=int(yr[c0]), month_label=m, value_hkd_m=ws.cell(5, c0).value,
                     index=ws.cell(5, c0 + 1).value, yoy_pct_reported=ws.cell(5, c0 + 2).value))
m = pd.DataFrame(rows)
m['provisional'] = m.month_label.str.contains(r'\sp$')
m['date'] = pd.to_datetime(m.year.astype(str) + '-' + m.month_label.str.replace(r'\sp$', '', regex=True) + '-01', format='%Y-%b-%d')
m = m.sort_values('date').reset_index(drop=True)
m['yoy_pct_recalc'] = (m.value_hkd_m / m.value_hkd_m.shift(12) - 1) * 100
# prior-year value: actual where the file has it, otherwise implied from reported yoy (2022 months)
m['prior_year_value'] = m.value_hkd_m.shift(12)
m['prior_source'] = np.where(m.prior_year_value.notna(), 'in_file', 'implied_from_reported_yoy')
imp = m.value_hkd_m / (1 + m.yoy_pct_reported / 100)
m['prior_year_value'] = m.prior_year_value.fillna(imp)
m['fy_start'] = np.where(m.date.dt.month >= 4, m.date.dt.year, m.date.dt.year - 1)
m['fq'] = ((m.date.dt.month - 4) % 12) // 3 + 1
m['fq_label'] = m.fq.astype(str) + 'Q' + m.fy_start.astype(str)
m['month_in_fq'] = ((m.date.dt.month - 4) % 3) + 1
m[['date', 'year', 'month_label', 'value_hkd_m', 'index', 'yoy_pct_reported', 'yoy_pct_recalc', 'provisional',
   'prior_year_value', 'prior_source', 'fq_label', 'month_in_fq']].to_csv(f'{OUT}/csd_jewellery_monthly_tidy.csv', index=False, date_format='%Y-%m-%d')

# ---------- 2. Fiscal-quarter yoy (3-month and first-two-month) ----------
q = []
for lab, g in m.groupby('fq_label', sort=False):
    g = g.sort_values('date')
    r = dict(fq_label=lab, start=g.date.min(), n_months=len(g), has_derived_prior=(g.prior_source != 'in_file').any(),
             last_month_provisional=bool(g.provisional.iloc[-1]))
    r['csd_yoy_3m'] = (g.value_hkd_m.sum() / g.prior_year_value.sum() - 1) * 100 if len(g) == 3 else np.nan
    g2 = g[g.month_in_fq <= 2]
    r['csd_yoy_2m'] = (g2.value_hkd_m.sum() / g2.prior_year_value.sum() - 1) * 100 if len(g2) == 2 else np.nan
    q.append(r)
q = pd.DataFrame(q)

# ---------- 3. Luk Fook SSSG tidy (read straight from the sheet) ----------
w = openpyxl.load_workbook(LF, data_only=True).active
recs = []
for c in range(2, w.max_column + 1):
    lab = w.cell(12, c).value
    if lab is None: continue
    raw = w.cell(13, c).value
    tag = w.cell(10, c).value
    note = w.cell(11, c).value
    val = raw * 100 if isinstance(raw, (int, float)) else np.nan
    status = 'ok' if isinstance(raw, (int, float)) else ('NA_text' if isinstance(raw, str) else 'blank')
    recs.append(dict(fq_label=lab, sssg_pct=val, raw_cell=raw, status=status, overseas_tag_row10=tag, note_row11=note))
s = pd.DataFrame(recs)
s.to_csv(f'{OUT}/lukfook_sssg_tidy.csv', index=False)

# ---------- 4. Gold: WGC quarterly average LBMA Gold Price PM (US$/oz), calendar quarters ----------
gold_usd = {  # (calendar year, calendar quarter): avg US$/oz; source: World Gold Council Gold Demand Trends tables (ICE Benchmark Administration data)
    (2022, 2): 1870.6, (2022, 3): 1728.9, (2022, 4): 1725.9,
    (2023, 1): 1889.9, (2023, 2): 1975.9, (2023, 3): 1928.5, (2023, 4): 1971.5,
    (2024, 1): 2069.8, (2024, 2): 2338.2, (2024, 3): 2474.3, (2024, 4): 2663.4,
    (2025, 1): 2859.6, (2025, 2): 3280.4, (2025, 3): 3456.5, (2025, 4): 4135.2,
    (2026, 1): 4872.9, (2026, 2): 4506.3,
}  # (2026,3) NOT AVAILABLE: LBMA Q3-2026 report gives low/high/start/end but no quarterly average
FX_BASE = 7.80  # USD/HKD peg midpoint placeholder (band 7.75-7.85). Constant FX leaves yoy identical to USD yoy.
gr = []
for (cy, cq), p in gold_usd.items():
    fy = cy if cq >= 2 else cy - 1
    fq = {2: 1, 3: 2, 4: 3, 1: 4}[cq]
    gr.append(dict(cal_year=cy, cal_quarter=cq, fq_label=f'{fq}Q{fy}', gold_usd_oz=p, usdhkd=FX_BASE, gold_hkd_oz=p * FX_BASE,
                   source='WGC Gold Demand Trends (quarterly avg LBMA Gold Price PM)'))
gdf = pd.DataFrame(gr).sort_values(['cal_year', 'cal_quarter']).reset_index(drop=True)
gdf['gold_hkd_yoy_pct'] = (gdf.gold_hkd_oz / gdf.gold_hkd_oz.shift(4) - 1) * 100
gdf['gold_usd_yoy_pct'] = (gdf.gold_usd_oz / gdf.gold_usd_oz.shift(4) - 1) * 100
gdf.to_csv(f'{OUT}/gold_lbma_quarterly_sourced.csv', index=False)

# ---------- 5. Merged panel (all fiscal quarters through 2Q2026) ----------
def qstart(lab):
    n, y = int(lab[0]), int(lab[2:]); return pd.Timestamp(y, 4, 1) + pd.DateOffset(months=3 * (n - 1))
labs = [f'{n}Q{y}' for y in range(2022, 2027) for n in range(1, 5)][:-2]  # through 2Q2026
P = pd.DataFrame({'fq_label': labs})
P['start'] = P.fq_label.map(qstart)
P['period'] = P.start.dt.strftime('%b%y') + '-' + (P.start + pd.DateOffset(months=2)).dt.strftime('%b%y')
P = P.merge(q.drop(columns='start'), on='fq_label', how='left').merge(s[['fq_label', 'sssg_pct', 'status', 'overseas_tag_row10']], on='fq_label', how='left') \
     .merge(gdf[['fq_label', 'gold_hkd_yoy_pct', 'gold_usd_yoy_pct']], on='fq_label', how='left')
# COVID-contamination: any month in the quarter OR its prior-year base window falls on/before Feb-2023
# (border fully reopened 6 Feb 2023, mask mandate ended 1 Mar 2023)  <=>  quarter starts on/before Feb-2024
P['covid_contaminated'] = P.start <= pd.Timestamp(2024, 2, 1)
P.to_csv(f'{OUT}/quarterly_panel.csv', index=False, date_format='%Y-%m-%d')
pd.set_option('display.width', 250); pd.set_option('display.max_columns', 30)
print(P.round(1).to_string())
print(m[['date', 'value_hkd_m', 'yoy_pct_reported', 'yoy_pct_recalc', 'provisional']].tail(4))
